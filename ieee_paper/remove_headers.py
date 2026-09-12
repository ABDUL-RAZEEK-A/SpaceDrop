from docx import Document

doc = Document('ieee_paper/SpaceDrop_IEEE_Paper.docx')
for section in doc.sections:
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    # Clear the header
    for paragraph in section.header.paragraphs:
        paragraph.text = ""
    # Clear the footer
    for paragraph in section.footer.paragraphs:
        paragraph.text = ""
    # Also explicitly clear them if they are unlinked
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

doc.save('ieee_paper/SpaceDrop_IEEE_Paper.docx')
print("Headers and footers removed")
